import streamlit as st
from docx import Document
from io import BytesIO, StringIO
import pandas as pd
from modules.common import get_client

def generate_article(client, topic):
    try:
        response = client.chat.completions.create(
            model="openai/gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "Eres un investigador experto en educacion con enfoque STEAM (Ciencia, Tecnologia, Ingenieria, Arte y Matematicas), con experiencia en diseno de propuestas pedagogicas e investigacion educativa. Redactas documentos academicos estructurados, claros, coherentes y aplicables en contextos reales."},
                {"role": "user", "content": f"""
Elabora un articulo academico aplicado a la educacion con enfoque STEAM sobre el siguiente tema: {topic}

El documento debe estar bien estructurado y redactado en un lenguaje academico claro, incluyendo las siguientes secciones:

1. Titulo
2. Resumen (maximo 150 palabras)
3. Introduccion (planteamiento del problema y contexto educativo)
4. Justificacion (importancia del tema en educacion y en el enfoque STEAM)
5. Objetivo general y objetivos especificos
6. Marco teorico (conceptos clave relacionados con STEAM y el tema)
7. Metodologia (tipo de enfoque, poblacion, instrumentos y procedimiento)
8. Propuesta o intervencion educativa con enfoque STEAM:
- Integracion de Ciencia, Tecnologia, Ingenieria, Arte y Matematicas
- Actividades practicas basadas en un problema real
- Uso de datos o analisis dentro de la propuesta
9. Resultados esperados
10. Discusion
11. Conclusiones

Requisitos importantes:
- El contenido debe ser coherente, aplicable y orientado a un contexto educativo real
- Debe evidenciar claramente el enfoque STEAM de forma integrada, no superficial
- Incluir ejemplos de actividades o aplicaciones practicas cuando sea posible
- Mantener un tono academico, pero comprensible para docentes

No incluyas explicaciones externas ni comentarios adicionales, solo el contenido del articulo.
El documento debe tener entre 800 y 1200 palabras en total.
Cada seccion debe ser breve, clara y bien estructurada.
"""}
            ],
            max_tokens=3000
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        err = str(e)
        if "401" in err and "User not found" in err:
            st.error("🔑 API Key inválida (401 User not found). Crea una nueva en https://openrouter.ai/keys y actualiza OPENROUTER_API_KEY en Render → Environment y en `.streamlit/secrets.toml`.")
        else:
            st.error(f"Error al generar articulo: {str(e)}")
        return None

def generate_code(client, description):
    try:
        response = client.chat.completions.create(
            model="openai/gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "Eres un experto en analisis de datos educativos con Python, con experiencia en ensenanza y enfoque STEAM. Generas codigo claro, funcional y orientado a la interpretacion de datos en contextos educativos reales."},
                {"role": "user", "content": f"""
Genera un script en Python para analizar datos relacionados con: {description}

El script debe estar orientado a un contexto educativo con enfoque STEAM y cumplir con lo siguiente:

1. Cargar un dataset (usar pandas)
2. Mostrar las primeras filas del dataset
3. Realizar analisis descriptivo basico (promedios, conteos, agrupaciones, etc.)
4. Aplicar al menos una transformacion o filtrado de datos
5. Generar resultados claros que permitan interpretar la informacion
6. Incluir al menos una visualizacion simple (por ejemplo, grafico de barras o lineas usando matplotlib o seaborn)
7. Ser entendible para estudiantes o docentes con conocimientos basicos de Python

Requisitos importantes:
- El codigo debe ser limpio, organizado y funcional
- No incluir explicaciones ni comentarios fuera del codigo
- Usar nombres de variables claros
- Enfocar el analisis en la comprension de datos dentro de un contexto educativo o investigativo

Devuelve unicamente el codigo en Python.
"""}
            ],
            temperature=0.2,
        )
        code = response.choices[0].message.content.strip()
        # Extrae bloque ```python si existe, sino devuelve tal cual (fix del bug original que duplicaba lineas)
        if "```" in code:
            lines = []
            in_block = False
            for line in code.split("\n"):
                if line.strip().startswith("```python"):
                    in_block = True
                    continue
                elif line.strip().startswith("```") and in_block:
                    in_block = False
                    continue
                elif in_block:
                    lines.append(line)
                elif not line.strip().startswith("```") and not in_block and "```" not in line:
                    # fuera de bloque sin markdown, ignorar hasta encontrar bloque
                    pass
            if lines:
                return "\n".join(lines).strip()
        return code.strip()
    except Exception as e:
        err = str(e)
        if "401" in err and "User not found" in err:
            st.error("🔑 API Key inválida (401). Crea una nueva en https://openrouter.ai/keys y actualiza Render → Environment.")
        else:
            st.error(f"Error al generar codigo: {str(e)}")
        return None

def generate_data_table(client, description):
    try:
        response = client.chat.completions.create(
            model="openai/gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "Eres un experto en diseno de datasets para investigacion educativa con enfoque STEAM (Ciencia, Tecnologia, Ingenieria, Arte y Matematicas). Generas datos realistas, coherentes y utiles para analisis e interpretacion en contextos educativos."},
                {"role": "user", "content": f"""
Genera un dataset en formato CSV relacionado con el siguiente tema: {description}

El dataset debe estar orientado a un contexto educativo o investigativo con enfoque STEAM y cumplir con lo siguiente:

1. Representar un problema real o situacion educativa (por ejemplo: rendimiento academico, uso de tecnologia, variables ambientales, etc.)
2. Incluir entre 10 y 20 filas de datos
3. Tener entre 4 y 6 columnas con nombres claros y significativos
4. Incluir al menos:
- una variable categorica (ej: grupo, curso, categoria)
- una variable numerica (ej: puntaje, cantidad, nivel)
5. Permitir analisis con Python (promedios, agrupaciones, comparaciones)
6. Ser coherente y realista (no valores aleatorios sin sentido)

Requisitos importantes:
- Usa comas como separador
- Incluye encabezados
- No incluyas texto adicional, solo el CSV limpio
- No uses formato markdown ni explicaciones

El dataset debe ser util para analisis de datos en un contexto educativo con enfoque STEAM.
"""}
            ],
            temperature=0.3,
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        err = str(e)
        if "401" in err and "User not found" in err:
            st.error("🔑 API Key inválida (401). Crea una nueva en https://openrouter.ai/keys y actualiza Render.")
        else:
            st.error(f"Error al generar tabla: {str(e)}")
        return None

def create_word_doc(article):
    doc = Document()
    doc.add_heading(text='Articulo Generado', level=0)
    doc.add_paragraph(article)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def create_excel(data):
    df = pd.read_csv(StringIO(data))
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Datos Generados')
    buffer.seek(0)
    return buffer, df

def render():
    st.title("STEAM Lab AI")
    client = get_client()
    if client is None:
        st.warning("Falta OPENROUTER_API_KEY en .streamlit/secrets.toml — las generaciones fallaran hasta configurarla.")

    # Fix bug original: el selectbox y los if deben coincidir exactamente
    section = st.selectbox(
        'Elige una seccion',
        ['Articulos de investigacion con enfoque STEAM','Scripts de investigacion con enfoque STEAM','Datasets para investigacion con enfoque STEAM'],
        key="steam_section"
    )

    if section == 'Articulos de investigacion con enfoque STEAM':
        st.header('Articulos de investigacion con enfoque STEAM')
        topic = st.text_input('Ingresa un tema para el articulo:',
                                placeholder='Ejemplo: "Integracion de robotica educativa en el aula de primaria"',
                                key="steam_topic")
        if st.button('Generar Articulo', key="steam_btn_article"):
            if topic:
                if client is None:
                    st.error("Configura OPENROUTER_API_KEY primero.")
                else:
                    with st.spinner('Generando Articulo...'):
                        article = generate_article(client, topic)
                        if article:
                            st.success('Articulo generado exitosamente')
                            st.markdown("### Vista previa del articulo")
                            st.markdown(article)
                            st.download_button(
                                'Descargar como Word',
                                data=create_word_doc(article),
                                file_name='articulo.docx',
                                mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                                key="steam_dl_word"
                            )
            else:
                st.warning('Por favor ingresa un tema para generar el articulo.')

    elif section == 'Scripts de investigacion con enfoque STEAM':
        st.header('Scripts en python de investigacion con enfoque STEAM')
        description = st.text_input('Describe el analisis de datos que deseas realizar',
                                    placeholder='Ejemplo: "Analizar el rendimiento academico de estudiantes en funcion del uso de tecnologia educativa"',
                                    key="steam_desc_code")
        if st.button('Generar Script', key="steam_btn_code"):
            if description:
                if client is None:
                    st.error("Configura OPENROUTER_API_KEY primero.")
                else:
                    with st.spinner('Generando Script en python...'):
                        code = generate_code(client, description)
                        if code:
                            st.success('Script generado exitosamente')
                            st.code(code, language='python')
                            st.download_button(
                                'Descargar Script',
                                data=code,
                                file_name='script.py',
                                mime='text/plain',
                                key="steam_dl_code"
                            )
            else:
                st.warning('Por favor ingresa una descripcion para generar el script.')

    elif section == 'Datasets para investigacion con enfoque STEAM':
        st.header('Datasets para investigacion con enfoque STEAM')
        data_description = st.text_input('Describe el tipo de datos que necesitas para tu investigacion o actividad STEAM:',
                                        placeholder='Ejemplo: "Datos de rendimiento academico de estudiantes en funcion del uso de tecnologia educativa"',
                                        key="steam_desc_data")
        if st.button('Generar Dataset', key="steam_btn_data"):
            if data_description:
                if client is None:
                    st.error("Configura OPENROUTER_API_KEY primero.")
                else:
                    with st.spinner('Generando Dataset...'):
                        data = generate_data_table(client, data_description)
                        if data:
                            st.success('Dataset generado exitosamente')
                            try:
                                excel_file, df = create_excel(data)
                                if excel_file and df is not None:
                                    st.markdown("### Vista previa del Dataset")
                                    st.dataframe(df, use_container_width=True)
                                    st.download_button(
                                        'Descargar Dataset en Excel',
                                        data=excel_file,
                                        file_name='tabla_datos.xlsx',
                                        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                        key="steam_dl_excel"
                                    )
                            except Exception as e:
                                st.error(f"Error al procesar CSV generado: {e}")
                                st.code(data)
            else:
                st.warning('Por favor ingresa una descripcion para generar el dataset.')
